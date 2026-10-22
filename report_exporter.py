
"""
個人財務管理系統
第 32 階段：PDF / Excel 報告匯出

修正版：
1. 修正 PDF 中文亂碼
2. 自動尋找 Windows 中文字型
3. 保留六個 Excel 工作表
4. 保留 Excel 支出分類圖表
5. PDF 支援自動換頁
"""

from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle,
)
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
)
from openpyxl.chart import (
    BarChart,
    Reference,
)


# ==========================================
# 1. 基本設定
# ==========================================

NAVY = "172C46"
GREEN = "16866B"
WHITE = "FFFFFF"

MONEY_FORMAT = '#,##0.00;[Red](#,##0.00)'

PDF_FONT_NAME = "FinanceChinese"


def format_money(value):
    return f"NT$ {float(value):,.0f}"


def safe_text(value):
    """避免特殊符號影響 PDF 文字排版。"""
    if value is None:
        return ""

    return escape(str(value))


# ==========================================
# 2. PDF 中文字型設定
# ==========================================

def register_chinese_font():
    """
    尋找支援繁體中文的 TrueType 字型。

    優先順序：
    1. 專案 fonts 資料夾
    2. Windows 微軟正黑體
    3. Windows 微軟雅黑
    4. Windows 細明體
    """

    if PDF_FONT_NAME in pdfmetrics.getRegisteredFontNames():
        return PDF_FONT_NAME

    project_dir = Path(__file__).resolve().parent

    candidates = [
        # 使用者自行放入的中文字型
        project_dir / "fonts" / "NotoSansTC-Regular.ttf",
        project_dir / "fonts" / "NotoSansTC.ttf",
        project_dir / "fonts" / "TaipeiSansTCBeta-Regular.ttf",

        # Windows 系統中文字型
        Path("C:/Windows/Fonts/msjh.ttc"),
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/mingliu.ttc"),
    ]

    errors = []

    for font_path in candidates:
        if not font_path.is_file():
            continue

        try:
            pdfmetrics.registerFont(
                TTFont(
                    PDF_FONT_NAME,
                    str(font_path),
                    subfontIndex=0,
                )
            )

            return PDF_FONT_NAME

        except Exception as error:
            errors.append(
                f"{font_path.name}: {error}"
            )

    details = (
        "\n".join(errors)
        if errors
        else "找不到可使用的中文字型。"
    )

    raise RuntimeError(
        "PDF 中文字型載入失敗。\n"
        "請在專案的 fonts 資料夾放入 "
        "NotoSansTC-Regular.ttf。\n"
        f"詳細資訊：\n{details}"
    )


# ==========================================
# 3. PDF 報告匯出
# ==========================================

def export_pdf(report):
    """
    將財務報告轉換成 PDF。

    回傳：
        bytes
    """

    font_name = register_chinese_font()

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=17 * mm,
        rightMargin=17 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=f"{report['month']} 財務報告",
        author="Personal Finance Analysis",
    )

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name="FinanceTitle",
            fontName=font_name,
            fontSize=20,
            leading=30,
            textColor=colors.HexColor("#172C46"),
            alignment=TA_CENTER,
            spaceAfter=12,
            wordWrap="CJK",
        )
    )

    styles.add(
        ParagraphStyle(
            name="FinanceHeading",
            fontName=font_name,
            fontSize=13,
            leading=20,
            textColor=colors.HexColor("#16866B"),
            spaceBefore=14,
            spaceAfter=9,
            wordWrap="CJK",
            keepWithNext=True,
        )
    )

    styles.add(
        ParagraphStyle(
            name="FinanceBody",
            fontName=font_name,
            fontSize=9,
            leading=16,
            textColor=colors.HexColor("#334155"),
            spaceAfter=6,
            wordWrap="CJK",
        )
    )

    styles.add(
        ParagraphStyle(
            name="FinanceCell",
            fontName=font_name,
            fontSize=8,
            leading=13,
            textColor=colors.HexColor("#334155"),
            wordWrap="CJK",
        )
    )

    styles.add(
        ParagraphStyle(
            name="FinanceHeaderCell",
            fontName=font_name,
            fontSize=8,
            leading=13,
            textColor=colors.white,
            wordWrap="CJK",
        )
    )

    story = []

    def add_text(text):
        story.append(
            Paragraph(
                safe_text(text),
                styles["FinanceBody"],
            )
        )

    def add_heading(text):
        story.append(
            Paragraph(
                safe_text(text),
                styles["FinanceHeading"],
            )
        )

    def add_table(headers, rows, widths):
        data = []

        header_row = [
            Paragraph(
                safe_text(value),
                styles["FinanceHeaderCell"],
            )
            for value in headers
        ]

        data.append(header_row)

        for row in rows:
            data.append([
                Paragraph(
                    safe_text(value),
                    styles["FinanceCell"],
                )
                for value in row
            ])

        table = Table(
            data,
            colWidths=widths,
            repeatRows=1,
            hAlign="LEFT",
            splitByRow=1,
        )

        table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#172C46"),
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F1F5F9"),
                    ],
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.HexColor("#DCE4ED"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
            ])
        )

        story.append(table)
        story.append(Spacer(1, 8))

    # ----------------------------------
    # 報告標題
    # ----------------------------------

    story.append(
        Paragraph(
            safe_text(
                f"{report['month']} 財務報告"
            ),
            styles["FinanceTitle"],
        )
    )

    add_text("Personal Finance Analysis")

    # ----------------------------------
    # 一、財務摘要
    # ----------------------------------

    add_heading("一、財務摘要")

    rate = report["savings_rate"]

    rate_text = (
        f"{rate:.1f}%"
        if rate is not None
        else "無法計算（本月無收入）"
    )

    add_table(
        ["項目", "數值"],
        [
            [
                "本月收入",
                format_money(report["income"]),
            ],
            [
                "本月支出",
                format_money(report["expense"]),
            ],
            [
                "本月結餘",
                format_money(report["balance"]),
            ],
            [
                "儲蓄率",
                rate_text,
            ],
            [
                "記帳筆數",
                report["record_count"],
            ],
        ],
        [
            65 * mm,
            110 * mm,
        ],
    )

    # ----------------------------------
    # 二、支出分類
    # ----------------------------------

    add_heading("二、支出分類分析")

    if report["categories"]:
        rows = []

        for row in report["categories"]:
            rows.append([
                row["category"],
                format_money(row["amount"]),
                f"{row['percentage']:.1f}%",
            ])

        add_table(
            [
                "分類",
                "支出金額",
                "占比",
            ],
            rows,
            [
                65 * mm,
                70 * mm,
                40 * mm,
            ],
        )

    else:
        add_text("本月沒有支出紀錄。")

    # ----------------------------------
    # 三、預算執行
    # ----------------------------------

    add_heading("三、預算執行情況")

    if report["budgets"]:
        rows = []

        for row in report["budgets"]:
            rows.append([
                row["category"],
                format_money(row["budget"]),
                format_money(row["spent"]),
                format_money(row["remaining"]),
            ])

        add_table(
            [
                "分類",
                "預算",
                "已使用",
                "剩餘",
            ],
            rows,
            [
                40 * mm,
                45 * mm,
                45 * mm,
                45 * mm,
            ],
        )

    else:
        add_text("本月尚未設定分類預算。")

    # ----------------------------------
    # 四、儲蓄目標
    # ----------------------------------

    add_heading("四、儲蓄目標進度")

    if report["savings_goals"]:
        rows = []

        for row in report["savings_goals"]:
            rows.append([
                row["name"],
                format_money(row["target"]),
                format_money(row["saved"]),
                f"{row['percentage']:.1f}%",
            ])

        add_table(
            [
                "目標名稱",
                "目標金額",
                "已存金額",
                "完成進度",
            ],
            rows,
            [
                50 * mm,
                50 * mm,
                45 * mm,
                30 * mm,
            ],
        )

    else:
        add_text("本月沒有儲蓄目標。")

    # ----------------------------------
    # 五、智慧財務建議
    # ----------------------------------

    add_heading("五、智慧財務建議")

    if report["insights"]:
        for index, message in enumerate(
            report["insights"],
            start=1,
        ):
            add_text(
                f"{index}. {message}"
            )

    else:
        add_text("目前沒有財務建議。")

    # ----------------------------------
    # 六、記帳明細
    # ----------------------------------

    add_heading("六、本月記帳明細")

    if report["transactions"]:
        rows = []

        for row in report["transactions"]:
            rows.append([
                row["date"],
                row["type"],
                row["category"],
                row["item"],
                format_money(row["amount"]),
            ])

        add_table(
            [
                "日期",
                "類型",
                "分類",
                "項目",
                "金額",
            ],
            rows,
            [
                29 * mm,
                20 * mm,
                33 * mm,
                48 * mm,
                45 * mm,
            ],
        )

    else:
        add_text("本月沒有記帳明細。")

    # ----------------------------------
    # 報告備註
    # ----------------------------------

    story.append(Spacer(1, 12))

    add_text(
        "報告說明：本報告根據記帳資料、"
        "分類預算與儲蓄目標自動產生。"
        "儲蓄目標存款獨立統計，"
        "不重複計入一般支出。"
    )

    # ----------------------------------
    # 頁碼
    # ----------------------------------

    def add_page_number(canvas, doc):
        canvas.saveState()

        canvas.setFont(
            font_name,
            8,
        )

        canvas.setFillColor(
            colors.HexColor("#64748B")
        )

        canvas.drawRightString(
            A4[0] - 17 * mm,
            10 * mm,
            f"第 {doc.page} 頁",
        )

        canvas.restoreState()

    document.build(
        story,
        onFirstPage=add_page_number,
        onLaterPages=add_page_number,
    )

    buffer.seek(0)

    return buffer.getvalue()


# ==========================================
# 4. Excel 報告匯出
# ==========================================

def export_excel(report):
    """
    將財務報告轉換成 Excel。

    回傳：
        bytes
    """

    wb = Workbook()

    wb.remove(wb.active)

    # ----------------------------------
    # 建立工作表共用函式
    # ----------------------------------

    def create_sheet(
        name,
        headers,
        rows,
    ):
        ws = wb.create_sheet(name)

        ws.append(headers)

        for row in rows:
            ws.append(row)

        ws.freeze_panes = "A2"

        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.font = Font(
                bold=True,
                color=WHITE,
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=NAVY,
            )

            cell.alignment = Alignment(
                horizontal="center",
            )

        ws.auto_filter.ref = ws.dimensions

        for column in ws.columns:
            letter = column[0].column_letter

            max_length = max(
                len(str(cell.value or ""))
                for cell in column
            )

            ws.column_dimensions[
                letter
            ].width = min(
                max(max_length + 4, 14),
                42,
            )

        for row in ws.iter_rows(
            min_row=2
        ):
            for cell in row:
                if isinstance(
                    cell.value,
                    (int, float),
                ):
                    cell.number_format = (
                        MONEY_FORMAT
                    )

        return ws

    # ----------------------------------
    # 工作表 1：財務摘要
    # ----------------------------------

    summary_rows = [
        [
            "月份",
            report["month"],
        ],
        [
            "收入",
            report["income"],
        ],
        [
            "支出",
            report["expense"],
        ],
        [
            "結餘",
            report["balance"],
        ],
        [
            "儲蓄率 (%)",
            report["savings_rate"],
        ],
        [
            "記帳筆數",
            report["record_count"],
        ],
    ]

    summary_ws = create_sheet(
        "財務摘要",
        ["項目", "數值"],
        summary_rows,
    )

    # 月份、儲蓄率、筆數不是金額
    for row_number in [2, 6, 7]:
        summary_ws.cell(
            row_number,
            2,
        ).number_format = "General"

    if report["savings_rate"] is None:
        summary_ws["B6"] = (
            "無法計算（本月無收入）"
        )

    # ----------------------------------
    # 工作表 2：支出分類
    # ----------------------------------

    category_ws = create_sheet(
        "支出分類",
        [
            "分類",
            "支出金額",
            "占比 (%)",
        ],
        [
            [
                row["category"],
                row["amount"],
                row["percentage"],
            ]
            for row in report["categories"]
        ],
    )

    for cell in category_ws["C"][1:]:
        cell.number_format = "0.0"

    if report["categories"]:
        chart = BarChart()

        chart.title = "各分類支出"

        chart.y_axis.title = (
            "金額（NT$）"
        )

        chart.x_axis.title = "分類"

        data = Reference(
            category_ws,
            min_col=2,
            min_row=1,
            max_row=category_ws.max_row,
        )

        labels = Reference(
            category_ws,
            min_col=1,
            min_row=2,
            max_row=category_ws.max_row,
        )

        chart.add_data(
            data,
            titles_from_data=True,
        )

        chart.set_categories(labels)

        chart.width = 16
        chart.height = 9

        category_ws.add_chart(
            chart,
            "E2",
        )

    # ----------------------------------
    # 工作表 3：預算執行
    # ----------------------------------

    budget_ws = create_sheet(
        "預算執行",
        [
            "分類",
            "預算金額",
            "實際支出",
            "剩餘金額",
            "使用率 (%)",
        ],
        [
            [
                row["category"],
                row["budget"],
                row["spent"],
                row["remaining"],
                row["percentage"],
            ]
            for row in report["budgets"]
        ],
    )

    for cell in budget_ws["E"][1:]:
        cell.number_format = "0.0"

    # ----------------------------------
    # 工作表 4：儲蓄目標
    # ----------------------------------

    savings_ws = create_sheet(
        "儲蓄目標",
        [
            "目標名稱",
            "目標金額",
            "已存金額",
            "剩餘金額",
            "完成率 (%)",
            "截止日期",
        ],
        [
            [
                row["name"],
                row["target"],
                row["saved"],
                row["remaining"],
                row["percentage"],
                row["target_date"],
            ]
            for row in report["savings_goals"]
        ],
    )

    for cell in savings_ws["E"][1:]:
        cell.number_format = "0.0"

    # ----------------------------------
    # 工作表 5：財務建議
    # ----------------------------------

    advice_ws = create_sheet(
        "財務建議",
        [
            "編號",
            "建議內容",
        ],
        [
            [
                index,
                message,
            ]
            for index, message in enumerate(
                report["insights"],
                start=1,
            )
        ],
    )

    for cell in advice_ws["A"][1:]:
        cell.number_format = "0"

    advice_ws.column_dimensions["B"].width = 65

    # ----------------------------------
    # 工作表 6：記帳明細
    # ----------------------------------

    create_sheet(
        "記帳明細",
        [
            "日期",
            "類型",
            "分類",
            "項目",
            "金額",
        ],
        [
            [
                row["date"],
                row["type"],
                row["category"],
                row["item"],
                row["amount"],
            ]
            for row in report["transactions"]
        ],
    )

    # ----------------------------------
    # 輸出 Excel
    # ----------------------------------

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    return output.getvalue()
